"""
Shared APA-7th house-style helpers for the PhDAI-733 Residency Project reports.

Mirrors the author's APA-strict report generators (Analytical Methods Assignment 4 and
the PhDAI-733 Hands-On projects): Georgia 11 pt, double-spaced, justified body with
0.5-in first-line indents, black APA headings, top-right page numbers, APA tables with
horizontal rules only, APA figures with bold number / italic title / "Note." beneath,
hanging-indent references. The report generator imports these helpers so every
document in this course shares one identity.
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Georgia"
SIZE = 11
BLACK = RGBColor(0, 0, 0)
LINE = 2.0
PAGE_TEXT_WIDTH = 6.5

AUTHORS = [("Yogananda Atreya Chinthaluri", "53651"), ("Anil Kumar Abburi", "53405")]
AFFIL = "School of Computer Science and Information Technology, University of the Cumberlands"
COURSE = "Python Application for Data Analytics in AI (PhDAI-733-M20)"
INSTRUCTOR = "Dr. Marwan Omar"
DATE = "October 31, 2026"

PCT = lambda x: f"{x * 100:.2f}%"
PCT1 = lambda x: f"{x * 100:.1f}%"


def new_doc():
    doc = Document()
    doc.core_properties.author = "; ".join(name for name, _ in AUTHORS)
    doc.core_properties.last_modified_by = doc.core_properties.author
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Inches(1.0)
        sec.left_margin = sec.right_margin = Inches(1.0)
    s = doc.styles["Normal"]
    s.font.name = FONT
    s.font.size = Pt(SIZE)
    s.font.color.rgb = BLACK
    pf = s.paragraph_format
    pf.line_spacing = LINE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    return doc


def _run(p, text, bold=False, italic=False, size=SIZE, color=BLACK, font=FONT):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = font
    r.font.size = Pt(size)
    r.font.color.rgb = color
    return r


def add_page_numbers(doc):
    for section in doc.sections:
        header = section.header
        header.is_linked_to_previous = False
        p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.line_spacing = 1.0
        for r in list(p.runs):
            r.text = ""
        run = p.add_run()
        run.font.name = FONT
        run.font.size = Pt(SIZE)
        run.font.color.rgb = BLACK
        b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
        i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = "PAGE"
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
        run._r.append(b); run._r.append(i); run._r.append(e)


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = LINE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p, text, bold=True)
    elif level == 2:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _run(p, text, bold=True)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _run(p, text, bold=True, italic=True)
    return p


def body(doc, text, indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = LINE
    p.paragraph_format.space_after = Pt(0)
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    _run(p, text)
    return p


def body_rich(doc, parts, indent=True):
    """
    Justified body paragraph assembled from (text, italic) fragments.

    Needed for APA mechanics: statistical symbols (r, n, M, SD) are set in italics,
    which a single plain-text run cannot express.
    """
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = LINE
    p.paragraph_format.space_after = Pt(0)
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    for text, italic in parts:
        _run(p, text, italic=italic)
    return p


def bullet(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.line_spacing = LINE
    p.paragraph_format.space_after = Pt(0)
    if bold_lead:
        _run(p, bold_lead, bold=True)
    _run(p, text)
    return p


def equation(doc, text):
    """Centered display equation in Cambria Math."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = LINE
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    _run(p, text, italic=True, font="Cambria Math")
    return p


def center_line(doc, text, bold=False, italic=False, size=SIZE, space_after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = LINE
    p.paragraph_format.space_after = Pt(space_after)
    _run(p, text, bold=bold, italic=italic, size=size)
    return p


# ── APA tables ────────────────────────────────────────────────────────────────
def _table_borders(tbl):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "6"); el.set(qn("w:color"), "000000")
        borders.append(el)
    for edge in ("left", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    tblPr.append(borders)


def _cell_bottom_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6"); bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    tcPr.append(borders)


def _cell_margins(cell, top=40, bottom=40, left=110, right=110):
    tcMar = OxmlElement("w:tcMar")
    for side, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        n = OxmlElement(f"w:{side}"); n.set(qn("w:w"), str(val)); n.set(qn("w:type"), "dxa"); tcMar.append(n)
    cell._tc.get_or_add_tcPr().append(tcMar)


def _row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true")
    trPr.append(cs)


def add_table(doc, number, title, headers, rows, col_widths=None, note=None, fs=9.5):
    p = doc.add_paragraph(); p.paragraph_format.line_spacing = LINE; p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True; p.paragraph_format.keep_together = True
    p.paragraph_format.space_before = Pt(6)
    _run(p, f"Table {number}", bold=True)
    p2 = doc.add_paragraph(); p2.paragraph_format.line_spacing = LINE; p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.keep_with_next = True; p2.paragraph_format.keep_together = True
    _run(p2, title, italic=True)

    if col_widths:
        tot = sum(col_widths); col_widths = [w * PAGE_TEXT_WIDTH / tot for w in col_widths]
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.autofit = False
    _table_borders(tbl)

    hdr = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        par = hdr[i].paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        par.paragraph_format.line_spacing = 1.0
        par.paragraph_format.space_after = Pt(0)
        _run(par, h, bold=True, size=fs)
        _cell_bottom_border(hdr[i]); _cell_margins(hdr[i])
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if col_widths:
            hdr[i].width = Inches(col_widths[i])
    for r_i, rd in enumerate(rows):
        cells = tbl.rows[r_i + 1].cells
        for c_i, val in enumerate(rd):
            cells[c_i].text = ""
            par = cells[c_i].paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            par.paragraph_format.line_spacing = 1.0
            par.paragraph_format.space_after = Pt(0)
            _run(par, str(val), size=fs)
            _cell_margins(cells[c_i]); cells[c_i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if col_widths:
                cells[c_i].width = Inches(col_widths[c_i])
    if col_widths:
        for col_idx, wd in enumerate(col_widths):
            for row in tbl.rows:
                row.cells[col_idx].width = Inches(wd)
    for row in tbl.rows:
        _row_cant_split(row)
    if note:
        pn = doc.add_paragraph(); pn.paragraph_format.line_spacing = LINE; pn.paragraph_format.space_after = Pt(0)
        pn.paragraph_format.keep_together = True
        _run(pn, "Note. ", italic=True, size=fs)
        _run(pn, note, size=fs)
    return tbl


def add_fig(doc, path, number, title, note=None, width=5.9):
    p = doc.add_paragraph(); p.paragraph_format.line_spacing = LINE; p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True; p.paragraph_format.keep_together = True
    p.paragraph_format.space_before = Pt(6)
    _run(p, f"Figure {number}", bold=True)
    p2 = doc.add_paragraph(); p2.paragraph_format.line_spacing = LINE; p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.keep_with_next = True; p2.paragraph_format.keep_together = True
    _run(p2, title, italic=True)
    pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.line_spacing = 1.0; pic.paragraph_format.space_before = Pt(4); pic.paragraph_format.space_after = Pt(2)
    pic.paragraph_format.keep_together = True
    if note:
        pic.paragraph_format.keep_with_next = True
    run = pic.add_run()
    run.add_picture(str(path), width=Inches(min(width, PAGE_TEXT_WIDTH)))
    if note:
        pn = doc.add_paragraph(); pn.paragraph_format.line_spacing = LINE; pn.paragraph_format.space_after = Pt(0)
        pn.paragraph_format.keep_together = True
        _run(pn, "Note. ", italic=True, size=10)
        _run(pn, note, size=10)


def ref_entry(doc, text):
    """Hanging-indent APA reference; '*...*' marks italics (journal/volume, titles)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.line_spacing = LINE
    p.paragraph_format.space_after = Pt(0)
    for i, seg in enumerate(text.split("*")):
        if seg:
            _run(p, seg, italic=i % 2 == 1)


def hyperlink(paragraph, url, text=None):
    """Append a clickable, underlined hyperlink run (python-docx has no public API)."""
    r_id = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, val in [("w:rFonts", None), ("w:color", "0563C1"), ("w:u", "single"),
                     ("w:sz", str(SIZE * 2))]:
        el = OxmlElement(tag)
        if tag == "w:rFonts":
            el.set(qn("w:ascii"), FONT); el.set(qn("w:hAnsi"), FONT)
        else:
            el.set(qn("w:val"), val)
        rpr.append(el)
    run.append(rpr)
    t = OxmlElement("w:t"); t.text = text or url
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def title_page(doc, title, subtitle=None):
    for _ in range(3):
        doc.add_paragraph().paragraph_format.line_spacing = LINE
    center_line(doc, title, bold=True)
    if subtitle:
        center_line(doc, subtitle, italic=True)
    center_line(doc, "")
    for name, student_id in AUTHORS:
        center_line(doc, f"{name} (Student ID {student_id})")
    center_line(doc, AFFIL)
    center_line(doc, COURSE)
    center_line(doc, INSTRUCTOR)
    center_line(doc, DATE)
    doc.add_page_break()


def references_section(doc, refs):
    doc.add_page_break()
    heading(doc, "References", 1)
    for r in sorted(refs, key=lambda e: e.replace("*", "").lower()):
        ref_entry(doc, r)


def verify(abstract, doc, refs):
    import re
    full_text = "\n".join(p.text for p in doc.paragraphs)
    surnames = [r.split(",")[0].strip() for r in sorted(refs)]
    uncited = [s for s in surnames if s not in full_text]
    n_words = len(re.findall(r"\S+", abstract))
    print(f"  abstract: {n_words} words (<=150) -> {'OK' if n_words <= 150 else 'TOO LONG'}")
    print(f"  references: {len(refs)} | uncited in body: {uncited if uncited else 'none'}")
    if uncited:
        raise SystemExit(f"ERROR: uncited references: {uncited}")
